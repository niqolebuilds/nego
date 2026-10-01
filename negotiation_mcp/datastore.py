"""Admin uploads of price data: parse, validate, preview, apply as a new version, roll back.

Nothing an admin uploads touches the live data until it is applied, and applying
never edits files in place: it copies the active data into a new version folder,
changes the copy, validates the whole book with the same rules as the engine, and
only then moves the ``ACTIVE`` pointer. Rolling back is moving the pointer again.
"""

from __future__ import annotations

import csv
import io
import shutil
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from . import appdb
from .engine import EngineError
from .pricebook import (
    ACTIVE_POINTER,
    DEFAULT_DATA_DIR,
    PRICE_HISTORY_COLUMNS,
    PRICE_HISTORY_REQUIRED,
    PRICE_INDEX_COLUMNS,
    REBATE_PROGRAM_COLUMNS,
    SAMPLE_MARKER,
    SKU_MASTER_COLUMNS,
    PriceBook,
    _parse_observation,
    active_data_dir,
)
from .settings import workspace

KINDS = {
    # kind: (file name, all columns, required columns, label)
    "price_history": ("price_history.csv", PRICE_HISTORY_COLUMNS, PRICE_HISTORY_REQUIRED, "Price list / records (POs, contracts, quotes)"),
    "sku_master": ("sku_master.csv", SKU_MASTER_COLUMNS, ["sku"], "SKU master (names, equivalence groups, single-source)"),
    "price_index": ("price_index.csv", PRICE_INDEX_COLUMNS, PRICE_INDEX_COLUMNS, "Price index (inflation)"),
    "rebate_programs": ("rebate_programs.csv", REBATE_PROGRAM_COLUMNS, REBATE_PROGRAM_COLUMNS, "Rebate programs"),
}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 200_000


def _versions_dir() -> Path:
    d = workspace() / "versions"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _staging_dir() -> Path:
    d = workspace() / "staging"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """CSV or XLSX (first sheet) to headers and rows of strings."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise EngineError(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        headers = [(h or "").strip() for h in (reader.fieldnames or [])]
        rows = [{(k or "").strip(): (v or "").strip() for k, v in r.items()} for r in reader]
    elif name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - dependency is in requirements
            raise EngineError("Excel support needs openpyxl; upload a CSV instead") from None
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise EngineError("That file isn't a readable .xlsx workbook") from None
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        first = next(it, None) or ()
        headers = [_cell(h) for h in first]
        rows = []
        for values in it:
            if values is None or all(v in (None, "") for v in values):
                continue
            rows.append({h: _cell(v) for h, v in zip(headers, values) if h})
        wb.close()
    else:
        raise EngineError("Upload a .csv or .xlsx file")
    if len(rows) > MAX_ROWS:
        raise EngineError(f"Too many rows ({len(rows):,}); the limit is {MAX_ROWS:,} per upload")
    return headers, rows


def _validate(kind: str, headers: list[str], rows: list[dict[str, str]]) -> dict:
    fname, columns, required, label = KINDS[kind]
    missing = [c for c in required if c not in headers]
    if missing:
        raise EngineError(f"Missing column(s) {missing}. Download the {label} template to see the expected headers.")
    if not rows:
        raise EngineError("The file has headers but no rows")
    errors: list[str] = []
    summary: dict[str, Any] = {"rows": len(rows), "unknown_columns": [h for h in headers if h and h not in columns]}
    if kind == "price_history":
        obs = []
        for i, r in enumerate(rows, start=2):
            try:
                obs.append(_parse_observation(r, fname, i))
            except EngineError as e:
                errors.append(str(e))
                if len(errors) >= 25:
                    break
        if obs:
            current = _safe_current_book()
            known_skus = set(current.skus) if current else set()
            known_vendors = {o.vendor for o in current.observations} if current else set()
            summary.update(
                date_from=min(o.date for o in obs).isoformat(),
                date_to=max(o.date for o in obs).isoformat(),
                by_source={s: sum(1 for o in obs if o.source == s) for s in ("po", "contract", "quote")},
                skus=len({o.sku for o in obs}),
                vendors=len({o.vendor for o in obs}),
                hospitals=len({o.hospital for o in obs}),
                new_skus=sorted({o.sku for o in obs} - known_skus)[:20],
                new_vendors=sorted({o.vendor for o in obs} - known_vendors)[:20],
            )
    else:
        for i, r in enumerate(rows, start=2):
            for c in required:
                if not r.get(c):
                    errors.append(f"{fname} row {i}: '{c}' is empty")
            if len(errors) >= 25:
                break
    summary["errors"] = errors
    summary["sample"] = rows[:5]
    return summary


def _safe_current_book() -> PriceBook | None:
    try:
        return PriceBook.load(active_data_dir())
    except EngineError:
        return None


def stage_upload(kind: str, filename: str, content: bytes, user_email: str) -> dict:
    """Parse and validate an upload, keep it in staging, and return the preview."""
    if kind not in KINDS:
        raise EngineError(f"Unknown upload type '{kind}'")
    headers, rows = read_table(filename, content)
    summary = _validate(kind, headers, rows)
    uid = uuid.uuid4().hex[:12]
    columns = KINDS[kind][1]
    with (_staging_dir() / f"{uid}.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in columns})
    status = "invalid" if summary["errors"] else "staged"
    safe_name = Path(filename).name[:200]
    appdb.add_upload(uid, user_email, kind, safe_name, len(rows), status, summary)
    appdb.audit(user_email, "upload.stage", {"upload": uid, "kind": kind, "file": safe_name, "rows": len(rows), "status": status})
    return {"id": uid, "kind": kind, "filename": safe_name, "status": status, **summary}


def _current_version_id() -> str | None:
    pointer = workspace() / ACTIVE_POINTER
    return pointer.read_text(encoding="utf-8").strip() or None if pointer.exists() else None


def apply_upload(uid: str, mode: str, user_email: str, note: str = "") -> dict:
    """Create a new data version with this upload added or replacing its file, then activate it."""
    up = appdb.get_upload(uid)
    if not up:
        raise EngineError("No such upload")
    if up["status"] != "staged":
        raise EngineError(f"This upload is {up['status']} and can't be applied")
    if mode not in ("add", "replace"):
        raise EngineError("Mode must be 'add' or 'replace'")
    kind = up["kind"]
    fname, columns, _, _ = KINDS[kind]
    base_dir = active_data_dir()
    base_id = _current_version_id() if base_dir != DEFAULT_DATA_DIR else "sample"
    vid = datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:4]
    new_dir = _versions_dir() / vid
    shutil.copytree(base_dir, new_dir)
    staged = _staging_dir() / f"{uid}.csv"
    target = new_dir / fname
    try:
        if mode == "replace" or not target.exists():
            shutil.copyfile(staged, target)
        else:
            with target.open(newline="", encoding="utf-8-sig") as f:
                existing = list(csv.DictReader(f))
            with staged.open(newline="", encoding="utf-8") as f:
                incoming = list(csv.DictReader(f))
            seen = {tuple(r.get(c, "") for c in columns) for r in existing}
            merged = existing + [r for r in incoming if tuple(r.get(c, "") for c in columns) not in seen]
            with target.open("w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
                w.writeheader()
                w.writerows({c: r.get(c, "") for c in columns} for r in merged)
        # Real price history replacing the sample means this version is no longer sample data.
        if kind == "price_history" and mode == "replace" and (new_dir / SAMPLE_MARKER).exists():
            (new_dir / SAMPLE_MARKER).unlink()
        book = PriceBook.load(new_dir)  # the whole book must still validate
    except Exception:
        shutil.rmtree(new_dir, ignore_errors=True)
        raise
    appdb.add_version(vid, user_email, base_id, note or f"{mode} {up['filename']}", len(book.observations), book.is_sample)
    _activate(vid)
    appdb.set_upload_status(uid, "applied", vid)
    staged.unlink(missing_ok=True)
    appdb.audit(user_email, "upload.apply", {"upload": uid, "mode": mode, "version": vid, "rows_now": len(book.observations)})
    return {"version": vid, "rows": len(book.observations), "is_sample": book.is_sample}


def discard_upload(uid: str, user_email: str) -> None:
    up = appdb.get_upload(uid)
    if not up or up["status"] not in ("staged", "invalid"):
        raise EngineError("Nothing to discard")
    (_staging_dir() / f"{uid}.csv").unlink(missing_ok=True)
    appdb.set_upload_status(uid, "discarded")
    appdb.audit(user_email, "upload.discard", {"upload": uid})


def _activate(vid: str | None) -> None:
    pointer = workspace() / ACTIVE_POINTER
    if vid is None:
        pointer.unlink(missing_ok=True)
        return
    if not (_versions_dir() / vid / "price_history.csv").exists():
        raise EngineError("That version no longer exists")
    tmp = pointer.with_suffix(".tmp")
    tmp.write_text(vid, encoding="utf-8")
    tmp.replace(pointer)


def activate_version(vid: str, user_email: str) -> dict:
    """Roll back or forward. ``sample`` switches to the bundled synthetic data."""
    if vid == "sample":
        _activate(None)
    else:
        if not any(v["id"] == vid for v in appdb.list_versions()):
            raise EngineError("No such version")
        _activate(vid)
    appdb.audit(user_email, "version.activate", {"version": vid})
    return versions()


def versions() -> dict:
    active = _current_version_id() if active_data_dir() != DEFAULT_DATA_DIR else None
    rows = appdb.list_versions()
    for r in rows:
        r["active"] = r["id"] == active
    return {"active": active or "sample", "versions": rows, "env_override": active_data_dir() not in (
        DEFAULT_DATA_DIR, *(_versions_dir() / r["id"] for r in rows))}


def template(kind: str) -> str:
    if kind not in KINDS:
        raise EngineError("Unknown template")
    buf = io.StringIO()
    csv.writer(buf).writerow(KINDS[kind][1])
    return buf.getvalue()
