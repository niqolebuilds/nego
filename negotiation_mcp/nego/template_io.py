"""Template_Nego in and out, keeping the workbook exactly as Siloam's team knows it.

``export_cycle`` fills the real template (title block, legend, section colours,
formulas, drop-downs and highlight rules stay as they are) and grows the formatted
area when a principal has more than the template's 150 rows. ``read_template`` reads
a filled workbook back (a principal's reply, or last cycle's file) by its headers,
and ``plan_import`` compares it with the cycle so a person sees every change before
it's applied.
"""

from __future__ import annotations

import io
import re
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any

from ..engine import EngineError
from . import model as M
from . import store
from .prepare import norm_name

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "Template_Nego.xlsx"
SHEET = "Form"
HEADER_ROW = 7
FIRST_ROW = 8
TEMPLATE_LAST_ROW = 157
LETTERS = [c.letter for c in M.COLUMNS]
BULAN = ("Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober",
         "November", "Desember")
BINDING_TEXT = {
    "Nett": "Binding: Nett Price (HNA after discount). Discount conversion applies if HNA increases during the contract period",
    "Disc": "Binding: Discount (% off HNA). The agreed discount applies to the HNA in force during the contract period",
}


def tanggal(iso: str) -> str:
    d = date.fromisoformat(iso)
    return f"{d.day:02d} {BULAN[d.month - 1]} {d.year}"


def _shift_formula(f: str, row: int, ppn: float) -> str:
    out = re.sub(r"(\$?[A-Z]{1,2})8\b", lambda m: f"{m.group(1)}{row}", f)
    if abs(ppn - M.DEFAULT_PPN) > 1e-12:
        out = out.replace("*1.11)", f"*{1 + ppn:g})")
    return out


def _extend_ranges(ws, last: int) -> None:
    """Grow validation and highlight ranges from row 157 to ``last``."""
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.worksheet.cell_range import MultiCellRange

    def grow(ranges: str) -> str:
        return " ".join(re.sub(rf"{TEMPLATE_LAST_ROW}$", str(last), r) for r in ranges.split())

    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(grow(str(dv.sqref)))
    old = ws.conditional_formatting
    new = ConditionalFormattingList()
    for cf in old:
        for rule in cf.rules:
            new.add(grow(str(cf.sqref)), rule)
    ws.conditional_formatting = new


def export_cycle(cycle: dict, items: list[dict], ppn: float = M.DEFAULT_PPN) -> bytes:
    import openpyxl

    if not TEMPLATE.exists():
        raise EngineError("Template_Nego.xlsx is missing from the app")
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET]
    p = cycle.get("principal") or {}
    ws["A1"] = f"Contract Period: {tanggal(cycle['contract_start'])}  -  {tanggal(cycle['contract_end'])}"
    ws["A2"] = BINDING_TEXT.get(cycle.get("binding") or "Nett", BINDING_TEXT["Nett"])
    ws["A3"] = f"Delivery Fee: {cycle.get('delivery_fee') or 'Free for all Siloam Hospitals units'}"
    ws["B4"] = p.get("name")
    ws["B5"] = p.get("distributor")

    last = max(TEMPLATE_LAST_ROW, FIRST_ROW + len(items) - 1)
    formulas = {c.letter: ws[f"{c.letter}{FIRST_ROW}"].value for c in M.COLUMNS if c.kind == "formula"}
    for r in range(FIRST_ROW, last + 1):
        if r > TEMPLATE_LAST_ROW:
            ws.row_dimensions[r].height = ws.row_dimensions[TEMPLATE_LAST_ROW].height
            for col in LETTERS:
                src, dst = ws[f"{col}{TEMPLATE_LAST_ROW}"], ws[f"{col}{r}"]
                dst._style = copy(src._style)
        for col, f in formulas.items():
            ws[f"{col}{r}"] = _shift_formula(f, r, ppn)
    if last > TEMPLATE_LAST_ROW:
        _extend_ranges(ws, last)

    for i, it in enumerate(items):
        r = FIRST_ROW + i
        for c in M.COLUMNS:
            if c.kind == "formula":
                continue
            v = it.get(c.field)
            ws[f"{c.letter}{r}"] = v if v not in ("",) else None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(h: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


HEADER_INDEX = {_norm(c.header): c for c in M.COLUMNS}


def read_template(content: bytes) -> dict:
    """Rows of a filled Template_Nego, by header. Formula columns are ignored (recomputed)."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    except Exception as e:  # noqa: BLE001
        raise EngineError(f"Can't read that Excel file: {e}") from None
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    hdr_row = None
    for r in range(1, 21):
        if any(_norm(ws.cell(r, c).value) == "erp code" for c in range(1, 30)):
            hdr_row = r
            break
    if hdr_row is None:
        raise EngineError("This doesn't look like Template_Nego: no 'ERP Code' header in the first 20 rows")
    cols: dict[int, M.Column] = {}
    for c in range(1, ws.max_column + 1):
        col = HEADER_INDEX.get(_norm(ws.cell(hdr_row, c).value))
        if col and col.kind != "formula":
            cols[c] = col
    if not any(col.field == "item_name" for col in cols.values()):
        raise EngineError("The template is missing the 'Item Name' column")

    def label(row: int) -> str | None:
        v = ws.cell(row, 2).value
        return str(v).strip() if v not in (None, "") else None

    out = {"principal": label(4), "distributor": label(5), "title": ws["A1"].value, "rows": [], "errors": []}
    for r in range(hdr_row + 1, ws.max_row + 1):
        vals = {col.field: ws.cell(r, c).value for c, col in cols.items()}
        if not vals.get("erp_code") and not vals.get("item_name"):
            continue
        row: dict[str, Any] = {"_row": r}
        for f, v in vals.items():
            if isinstance(v, str) and v.startswith("="):
                continue  # someone typed a formula into an input cell; ignore it
            try:
                row[f] = store.clean_value(f, v)
            except EngineError as e:
                out["errors"].append({"row": r, "field": f, "message": str(e)})
        out["rows"].append(row)
    wb.close()
    return out


def _match_key(row: dict) -> tuple[str | None, str | None]:
    code = str(row.get("erp_code") or "").strip().upper() or None
    return code, norm_name(row.get("item_name")) or None


def _differs(a: Any, b: Any) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) > 1e-9 * max(1.0, abs(float(a)))
    return a != b


def plan_import(cycle_items: list[dict], parsed: dict, fields: tuple[str, ...]) -> dict:
    """What applying a filled template would change. Blank cells never erase a value."""
    by_code = {str(i["erp_code"]).strip().upper(): i for i in cycle_items if i.get("erp_code")}
    by_name = {norm_name(i["item_name"]): i for i in cycle_items if i.get("item_name")}
    changes: list[dict] = []
    unmatched: list[dict] = []
    seen: set[int] = set()
    dup_rows: list[int] = []
    for row in parsed["rows"]:
        code, name = _match_key(row)
        item = by_code.get(code) if code else None
        if item is None and name:
            item = by_name.get(name)
        if item is None:
            unmatched.append({"row": row["_row"], "erp_code": row.get("erp_code"), "item_name": row.get("item_name")})
            continue
        if item["id"] in seen:
            dup_rows.append(row["_row"])
            continue
        seen.add(item["id"])
        for f in fields:
            if f not in row or row[f] is None:
                continue
            if _differs(item.get(f), row[f]):
                changes.append({"item_id": item["id"], "erp_code": item.get("erp_code"), "item_name": item.get("item_name"),
                                "field": f, "header": M.BY_FIELD[f].header, "old": item.get(f), "new": row[f]})
    counts: dict[str, int] = {}
    for c in changes:
        counts[c["header"]] = counts.get(c["header"], 0) + 1
    return {
        "rows_in_file": len(parsed["rows"]),
        "matched": len(seen),
        "items_changed": len({c["item_id"] for c in changes}),
        "changes": changes,
        "changes_by_column": counts,
        "unmatched": unmatched[:200],
        "unmatched_count": len(unmatched),
        "duplicate_rows": dup_rows[:50],
        "not_in_file": len(cycle_items) - len(seen),
        "errors": parsed["errors"][:200],
        "principal_in_file": parsed.get("principal"),
    }


def apply_plan(cid: int, plan: dict, by: str, via: str) -> int:
    per_item: dict[int, dict] = {}
    for c in plan["changes"]:
        per_item.setdefault(c["item_id"], {})[c["field"]] = c["new"]
    for item_id, ch in per_item.items():
        store.update_item(cid, item_id, ch, by, via=via)
    return len(per_item)
