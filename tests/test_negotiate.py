"""Phase 3: counter-offer suggestions, market benchmarks, Online Nego, escalation, the
submission package and the principal's encrypted documents."""

from __future__ import annotations

import io

import openpyxl
import pytest
from starlette.testclient import TestClient

from negotiation_mcp import appdb
from negotiation_mcp.nego import benchmark, model as M, negotiate, portal, service, store, vault

H = {"X-Requested-With": "nego"}


@pytest.fixture()
def ws(tmp_path, monkeypatch):
    monkeypatch.setenv("NEGO_HOME", str(tmp_path))
    monkeypatch.delenv("NEGO_DOC_KEY", raising=False)
    store.init()
    appdb.init()
    return tmp_path


BASE = {"item_status": "Active", "mou_qty": 10, "mou_hna": 100_000, "mou_disc": 0.1,
        "rfq_qty": 10, "rfq_hna": 110_000, "rfq_disc": 0.1}


# ---------------------------------------------------------------- counter offer

def test_counter_offer_holds_the_mou_price():
    s = negotiate.suggest_co(dict(BASE), None, 0.11, 0.25)
    co_price = M.unit_price(110_000, 10, s["co_disc"])
    mou_price = M.unit_price(100_000, 10, 0.1)
    assert co_price <= mou_price and co_price > mou_price * 0.998  # rounded up to 0.1%, never above the MOU
    assert "MOU" in s["note"]


def test_counter_offer_uses_the_lowest_reference_and_the_cap():
    bench = {"price_pp": 8_000.0, "source": "INAPROC"}
    s = negotiate.suggest_co(dict(BASE), bench, 0.11, 0.25)
    assert "INAPROC" in s["note"]
    assert M.unit_price(110_000, 10, s["co_disc"]) == pytest.approx(8_000, rel=0.002)
    capped = negotiate.suggest_co(dict(BASE), {"price_pp": 1_000.0, "source": "X"}, 0.11, 0.25)
    assert capped["co_disc"] == pytest.approx(0.35) and "Capped" in capped["note"]
    cheap = negotiate.suggest_co({**BASE, "rfq_hna": 90_000}, None, 0.11, 0.25)
    assert cheap["co_disc"] == 0.1 and "accept" in cheap["note"]
    assert negotiate.suggest_co({**BASE, "item_status": "Discontinue"}, None, 0.11, 0.25) is None


def make_cycle(n=4):
    p = store.create_principal({"name": "PT Nego Tiga", "contact_email": "a@b.example"}, "a@x")
    c = store.create_cycle(p["id"], {"contract_start": "2027-01-01", "contract_end": "2029-12-31"}, "a@x")
    rows = [{**BASE, "erp_code": f"T{i}", "item_name": f"Spuit {i}cc Terumo", "brand": "Terumo", "catalog_no": f"TER-{1000 + i}",
             "po_qty_12m": 100, "po_unit_text": "BOX@10", "last_po_price": 90_000} for i in range(1, n + 1)]
    store.add_items(c["id"], rows, "a@x")
    return c


def test_apply_co_fill_on_and_package(ws):
    c = make_cycle()
    plan = negotiate.plan_co(c["id"])
    assert plan["count"] == 4 and plan["co_impact"] <= 0
    assert negotiate.apply_co(c["id"], "a@x")["applied"] == 4
    it = store.items(c["id"])[0]
    assert it["co_note"] and it["co_disc"] > 0.1
    assert negotiate.plan_co(c["id"])["count"] == 0  # set ones are kept unless overwrite
    store.update_item(c["id"], it["id"], {"fb1_disc": 0.15, "price_reason": "Kurs"}, "p")
    assert negotiate.fill_on(c["id"], "a@x")["filled"] == 4
    got = {i["erp_code"]: i for i in store.items(c["id"])}
    assert got["T1"]["on_disc"] == 0.15 and got["T2"]["on_disc"] == got["T2"]["co_disc"]
    esc = negotiate.escalation(c["id"])
    assert esc["items_agreed"] == 4 and esc["needed"]  # T1 ends above the MOU
    name, data, info = negotiate.package_xlsx(c["id"])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Excel Confirmation", "BAK Draft", "Checks"]
    sh = wb["Excel Confirmation"]
    header_row = next(r for r in range(1, 15) if sh.cell(r, 2).value == "ERP Code")
    codes = [sh.cell(r, 2).value for r in range(header_row + 1, header_row + 5)]
    assert codes == ["T1", "T2", "T3", "T4"]  # A-Z by item name
    assert info["ready"]


def test_package_flags_missing_and_duplicates(ws):
    c = make_cycle(3)
    store.add_items(c["id"], [{**BASE, "erp_code": "T1", "item_name": "Duplicate"}], "a@x")
    chk = negotiate.package_checks(c["id"])
    assert not chk["ready"]
    assert any("no agreed" in p for p in chk["problems"]) and any("Duplicate ERP" in p for p in chk["problems"])
    with pytest.raises(Exception, match="No agreed prices"):
        negotiate.package_xlsx(c["id"])


# ---------------------------------------------------------------- benchmarks

BENCH = b"""Nama Produk,Merek,No Katalog,Satuan,Harga,Tanggal
Spuit 1cc Terumo,Terumo,TER-1001,BOX isi 10,85000,2026-08-01
TERUMO SPUIT 2CC,Terumo,,BOX isi 10,90000,2026-08-01
Spuit 3cc OneMed,OneMed,,BOX isi 10,70000,2026-08-01
Kasa Steril 16x16,OneMed,,PACK isi 10,50000,2026-08-01
"""


def test_benchmark_import_and_match(ws):
    c = make_cycle(3)
    r = benchmark.read_file("inaproc.csv", BENCH, "INAPROC", True, 0.11, "a@x")
    assert r["rows"] == 4
    assert store.benchmarks()[0]["price_pp"] == pytest.approx(8_500)
    res = benchmark.match_cycle(c["id"])
    ms = store.matches(c["id"])
    by = {(m["erp_code"], m["bm_name"]): m for m in ms}
    assert by[("T1", "Spuit 1cc Terumo")]["method"] == "ref" and by[("T1", "Spuit 1cc Terumo")]["status"] == "confirmed"
    assert by[("T2", "TERUMO SPUIT 2CC")]["status"] == "confirmed"  # same words, any order
    assert ("T3", "Kasa Steril 16x16") not in by  # a different product never matches
    assert all(not (m["erp_code"] == "T1" and m["bm_name"] == "TERUMO SPUIT 2CC" and m["status"] == "confirmed") for m in ms)
    assert res["items_matched"] >= 2
    best = store.best_benchmarks(c["id"])
    items = {i["erp_code"]: i["id"] for i in store.items(c["id"])}
    assert best[items["T1"]]["price_pp"] == pytest.approx(8_500)
    # the counter offer now uses the benchmark where it's the lowest reference
    s = negotiate.plan_co(c["id"])["suggestions"]
    t1 = next(x for x in s if x["item_id"] == items["T1"])
    assert "INAPROC" in t1["note"]
    # benchmark flag in Siloam's findings, never in the principal's view
    service.scan(c["id"])
    assert any(a["rule"] == "above_benchmark" for a in store.anomalies(c["id"], "open"))
    store.set_step(c["id"], "rfq", "a@x")
    view = portal.view_item(store.items(c["id"])[0], "rfq", 0.11, portal.thresholds())
    assert not {k for k in view if "bench" in k}


def test_decisions_are_kept_and_remembered(ws):
    c = make_cycle(3)
    benchmark.read_file("inaproc.csv", BENCH, "INAPROC", True, 0.11, "a@x")
    benchmark.match_cycle(c["id"])
    sug = store.matches(c["id"], "suggested")
    if sug:
        store.decide_match(c["id"], sug[0]["id"], "rejected", "a@x")
        benchmark.match_cycle(c["id"])
        assert any(m["id"] == sug[0]["id"] and m["status"] == "rejected" for m in store.matches(c["id"]))


# ---------------------------------------------------------------- documents

PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"


def test_documents_encrypted_and_required(ws):
    c = make_cycle(1)
    store.set_step(c["id"], "submission", "a@x")
    c = store.get_cycle(c["id"])
    s = portal.summary(c)
    assert not s["can_send"] and {o["issues"][0]["code"] for o in s["outstanding"]} == {"doc_missing"}
    with pytest.raises(Exception):
        portal.upload_document(c, "nib", "x.exe", b"MZ", "p")
    with pytest.raises(Exception):
        portal.upload_document(c, "nib", "fake.pdf", b"not a pdf", "p")
    portal.upload_document(c, "nib", "nib.pdf", PDF, "p")
    v = portal.upload_document(c, "npwp", "npwp.pdf", PDF, "p")
    assert v["missing_required"] == []
    d = store.documents(c["id"])[0]
    raw = (ws / "nego_docs" / str(c["id"]) / d["stored_name"]).read_bytes()
    assert b"%PDF" not in raw  # encrypted at rest
    assert vault.load(c["id"], d["stored_name"]) == PDF
    assert portal.summary(c)["can_send"]
    portal.submit(c, "p")
    with pytest.raises(Exception):
        portal.upload_document(store.get_cycle(c["id"]), "other", "late.pdf", PDF, "p")


def test_routes(ws):
    from negotiation_mcp.dashboard.app import create_app

    a = TestClient(create_app())
    a.post("/api/auth/signin", json={"email": "admin@example.com"}, headers=H)
    c = make_cycle(2)
    r = a.post("/api/admin/nego/benchmarks", files={"file": ("b.csv", BENCH)}, data={"source": "INAPROC", "incl_ppn": "1"}, headers=H)
    assert r.status_code == 200, r.text
    assert a.post(f"/api/admin/nego/cycles/{c['id']}/benchmarks/match", headers=H).status_code == 200
    assert a.get(f"/api/nego/cycles/{c['id']}/co-plan").json()["count"] == 2
    assert a.post(f"/api/admin/nego/cycles/{c['id']}/co", json={}, headers=H).json()["applied"] == 2
    assert a.post(f"/api/admin/nego/cycles/{c['id']}/on-fill", json={}, headers=H).json()["filled"] == 2
    assert a.patch(f"/api/admin/nego/cycles/{c['id']}", json={"meeting_at": "2026-11-02T10:00", "meeting_notes": "OK"}, headers=H).status_code == 200
    assert a.get(f"/api/nego/cycles/{c['id']}/package").json()["agreed"] == 2
    x = a.get(f"/api/admin/nego/cycles/{c['id']}/package.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    # principal sends documents; only admins can open them, and each view is logged
    store.set_step(c["id"], "submission", "a@x")
    link = a.post(f"/api/admin/nego/cycles/{c['id']}/links", json={}, headers=H).json()
    p = TestClient(a.app)
    resp = p.get(link["path"], follow_redirects=False)
    p.cookies.set("nego_principal", resp.cookies.get("nego_principal"))
    up = p.post("/api/p/documents", files={"file": ("nib.pdf", PDF)}, data={"doc_type": "nib"}, headers=H)
    assert up.status_code == 200, up.text
    docs = a.get(f"/api/admin/nego/cycles/{c['id']}/documents").json()["documents"]
    assert docs[0]["doc_type"] == "nib"
    dl = a.get(f"/api/admin/nego/cycles/{c['id']}/documents/{docs[0]['id']}")
    assert dl.content == PDF
    assert any(e["action"] == "document.view" for e in store.cycle_events(c["id"]))
    appdb.create_user("v@example.com", "Vee", "viewer", "admin@example.com")
    v = TestClient(a.app)
    v.post("/api/auth/signin", json={"email": "v@example.com"}, headers=H)
    assert v.get(f"/api/admin/nego/cycles/{c['id']}/documents/{docs[0]['id']}").status_code == 403
    assert v.post(f"/api/admin/nego/cycles/{c['id']}/co", json={}, headers=H).status_code == 403
